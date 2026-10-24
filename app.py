"""Phần mềm tính lương - Streamlit. Không yêu cầu tệp font ngoài."""
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
import unicodedata
import zipfile
import uuid

import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

st.set_page_config(page_title="Phần mềm tính tiền lương", page_icon="💰", layout="wide")


def lay_cau_hinh_supabase():
    """Lấy thông tin kết nối Supabase từ Streamlit Secrets."""
    try:
        cau_hinh = st.secrets["supabase"]
        url = str(cau_hinh["url"]).strip().rstrip("/")
        key = str(cau_hinh["secret_key"]).strip()

        if not url.startswith("https://") or not key:
            return None

        return {
            "url": url,
            "key": key
        }
    except (KeyError, FileNotFoundError, AttributeError):
        return None


def tai_pdf_len_supabase(pdf_bytes, duong_dan):
    """Tải PDF lên Storage riêng tư; không ghi đè tệp đã có."""
    from urllib.parse import quote
    cau_hinh = lay_cau_hinh_supabase()
    if cau_hinh is None:
        raise ValueError("Chưa cấu hình Supabase")
    if not isinstance(pdf_bytes, bytes) or not pdf_bytes.startswith(b"%PDF-"):
        raise ValueError("Dữ liệu PDF không hợp lệ")
    duong_dan = str(duong_dan).strip().lstrip("/")
    if not duong_dan or ".." in duong_dan.split("/"):
        raise ValueError("Đường dẫn PDF không hợp lệ")
    url = f"{cau_hinh['url']}/storage/v1/object/salary-pdfs/{quote(duong_dan, safe='/')}"
    try:
        response = requests.post(
            url,
            headers={
                "apikey": cau_hinh["key"],
                "Authorization": f"Bearer {cau_hinh['key']}",
                "Content-Type": "application/pdf",
                "x-upsert": "false",
            },
            data=pdf_bytes,
            timeout=30,
        )
    except requests.RequestException as exc:
        raise RuntimeError(
            "Unknown: Không rõ PDF đã được tải lên hay chưa; cần đối soát Storage."
        ) from exc
    if response.status_code not in (200, 201):
        raise RuntimeError(f"Không thể tải PDF lên Supabase: HTTP {response.status_code}")
    return duong_dan


def chuan_hoa_so_zalo(so_dien_thoai):
    """Chuẩn hóa số điện thoại Việt Nam để tìm kiếm Zalo."""
    if pd.isna(so_dien_thoai):
        return None

    so = str(so_dien_thoai).strip()

    # Loại bỏ khoảng trắng, dấu chấm và dấu gạch ngang
    so = re.sub(r"[\s.\-()]", "", so)

    # Chuyển đầu số quốc tế +84 hoặc 84 thành 0
    if so.startswith("+84"):
        so = "0" + so[3:]
    elif so.startswith("84") and len(so) == 11:
        so = "0" + so[2:]

    # Kiểm tra số di động Việt Nam gồm 10 chữ số
    if not re.fullmatch(r"0[35789]\d{8}", so):
        return None

    return so

def tao_cong_viec_zalo(batch_id, ho_ten, so_zalo, ky_luong, pdf_path):
    """Tạo công việc chờ Windows Agent, không tự gửi qua Zalo."""
    cau_hinh = lay_cau_hinh_supabase()
    if cau_hinh is None:
        raise ValueError("Chưa cấu hình Supabase")
    payload = {
        "batch_id": str(batch_id),
        "employee_name": str(ho_ten).strip(),
        "zalo_phone": str(so_zalo).strip(),
        "salary_period": str(ky_luong).strip(),
        "pdf_path": str(pdf_path).strip(),
        "status": "Pending",
    }
    try:
        response = requests.post(
            f"{cau_hinh['url']}/rest/v1/zalo_salary_jobs",
            headers={
                "apikey": cau_hinh["key"],
                "Authorization": f"Bearer {cau_hinh['key']}",
                "Content-Type": "application/json",
                "Prefer": "return=representation",
            },
            json=payload,
            timeout=20,
        )
    except requests.RequestException as exc:
        raise RuntimeError(
            "Unknown: Không xác định được Supabase đã tạo công việc hay chưa; "
            "cần đối soát trước khi thử lại."
        ) from exc
    if response.status_code == 409:
        raise RuntimeError(
            "Công việc đã tồn tại cho số Zalo và kỳ lương này; không tạo lại."
        )
    if response.status_code != 201:
        raise RuntimeError(f"Không thể tạo công việc Zalo: HTTP {response.status_code}")
    try:
        data = response.json()
        if not isinstance(data, list) or not data or not data[0].get("id"):
            raise ValueError("Phản hồi thiếu ID")
        return data[0]["id"]
    except (ValueError, TypeError, KeyError) as exc:
        raise RuntimeError(
            "Unknown: Supabase có thể đã tạo công việc nhưng phản hồi không đầy đủ; cần đối soát."
        ) from exc


MAX_NV = 500
COT_TIEN = ["Lương cơ bản", "Tiền thưởng", "Phụ cấp tính BHXH", "Phụ cấp không tính BHXH",
            "Tổng thu nhập", "Bảo hiểm người lao động đóng", "Giảm trừ bản thân",
            "Giảm trừ người phụ thuộc", "Thu nhập tính thuế", "Thuế TNCN", "Tiền thực lãnh"]
COT_NHAP = ["Họ và tên", "Số điện thoại Zalo", "Email cá nhân", "Ngày sinh", "Chức vụ", "Lương cơ bản", "Tiền thưởng",
            "Phụ cấp tính BHXH", "Phụ cấp không tính BHXH", "Số người phụ thuộc"]


def du_lieu_mac_dinh():
    return pd.DataFrame({
        "Họ và tên": [f"Nhân viên {i+1}" for i in range(MAX_NV)],
        "Số điện thoại Zalo": [""] * MAX_NV,
        "Email cá nhân": [""] * MAX_NV,
        "Ngày sinh": [""] * MAX_NV,
        "Chức vụ": [""] * MAX_NV,
        "Lương cơ bản": [15.0] * MAX_NV,
        "Tiền thưởng": [2.0] * MAX_NV,
        "Phụ cấp tính BHXH": [1.0] * MAX_NV,
        "Phụ cấp không tính BHXH": [0.5] * MAX_NV,
        "Số người phụ thuộc": [1] * MAX_NV,
    })


def tinh_thue_tncn(t):
    t = max(float(t), 0)
    if t <= 10:
        return t * 0.05
    if t <= 30:
        return t * 0.10 - 0.5
    if t <= 60:
        return t * 0.20 - 3.5
    if t <= 100:
        return t * 0.30 - 9.5
    return t * 0.35 - 14.5


def tinh_luong(nv):
    luong = float(nv["Lương cơ bản"])
    thuong = float(nv["Tiền thưởng"])
    pc_bh = float(nv["Phụ cấp tính BHXH"])
    pc_khac = float(nv["Phụ cấp không tính BHXH"])
    n = int(nv["Số người phụ thuộc"])
    tong = luong + thuong + pc_bh + pc_khac
    bao_hiem = (luong + pc_bh) * 0.105
    giam_tru = 15.5
    giam_tru_pt = 6.2 * n
    tinh_thue = max(0, tong - bao_hiem - giam_tru - giam_tru_pt)
    thue = tinh_thue_tncn(tinh_thue)
    return {"Tổng thu nhập": tong, "Bảo hiểm người lao động đóng": bao_hiem,
            "Giảm trừ bản thân": giam_tru, "Giảm trừ người phụ thuộc": giam_tru_pt,
            "Thu nhập tính thuế": tinh_thue, "Thuế TNCN": thue,
            "Tiền thực lãnh": tong - bao_hiem - thue}


def kiem_tra(df):
    loi = []
    for i, r in df.iterrows():
        dong = i + 1
        ten = str(r["Họ và tên"]).strip() if pd.notna(r["Họ và tên"]) else ""
        if not ten:
            loi.append(f"Dòng {dong}: thiếu họ tên")
        email = str(r["Email cá nhân"]).strip() if pd.notna(r["Email cá nhân"]) else ""
        if email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            loi.append(f"Dòng {dong}: email không hợp lệ")
        ngay = str(r["Ngày sinh"]).strip() if pd.notna(r["Ngày sinh"]) else ""
        if ngay:
            try:
                ngay_sinh = datetime.strptime(ngay, "%d/%m/%Y").date()
                if ngay_sinh > date.today():
                    loi.append(f"Dòng {dong}: ngày sinh ở tương lai")
            except ValueError:
                loi.append(f"Dòng {dong}: ngày sinh phải theo DD/MM/YYYY")
        for c in [
            "Lương cơ bản", "Tiền thưởng", "Phụ cấp tính BHXH",
            "Phụ cấp không tính BHXH", "Số người phụ thuộc",
        ]:
            try:
                v = float(r[c])
                if not (0 <= v < float("inf")):
                    raise ValueError
                if c == "Số người phụ thuộc" and (not v.is_integer() or v > 100):
                    raise ValueError
            except (ValueError, TypeError):
                loi.append(f"Dòng {dong}: {c} phải là số không âm hợp lệ")
    return loi


def tao_excel(df):
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD/MM/YYYY") as writer:
        df.to_excel(writer, sheet_name="Bang luong", index=False, startrow=2)
        df.to_excel(writer, sheet_name="Chi tiet", index=False)
        for name, hr in (("Bang luong", 3), ("Chi tiet", 1)):
            ws = writer.sheets[name]
            if name == "Bang luong":
                ws["A1"] = "BẢNG TỔNG HỢP LƯƠNG VÀ THUẾ TNCN"
                ws["A2"] = "Đơn vị tiền: triệu đồng/tháng"
            for cell in ws[hr]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[hr].height = 32
            ws.freeze_panes = f"A{hr+1}"
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(ws.max_column)}{ws.max_row}"
            for col_idx in range(1, ws.max_column + 1):
                title = str(ws.cell(hr, col_idx).value)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(title) + 4, 18), 34)
                if title in COT_TIEN:
                    for row in range(hr+1, ws.max_row+1):
                        ws.cell(row, col_idx).number_format = '#,##0.000;[Red]-#,##0.000;0.000'
        ws = writer.sheets["Bang luong"]
        end = ws.max_row
        total_row = end + 1
        ws.cell(total_row, 1, "TỔNG CỘNG")
        for col_idx in range(1, ws.max_column+1):
            title = ws.cell(3, col_idx).value
            if title in COT_TIEN:
                letter = get_column_letter(col_idx)
                ws.cell(total_row, col_idx, f"=SUM({letter}4:{letter}{end})")
                ws.cell(total_row, col_idx).number_format = '#,##0.000'
            cell = ws.cell(total_row, col_idx)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return buf.getvalue()


def ascii_pdf(s):
    """PDF dùng font chuẩn Helvetica: chuyển chữ có dấu sang không dấu."""
    s = str(s).replace("đ", "d").replace("Đ", "D")
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")


def tao_pdf(nv, ky_luong):
    """Không dùng font ngoài; nội dung PDF tiếng Việt KHÔNG DẤU."""
    from xml.sax.saxutils import escape
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    elements = [Paragraph("PHIEU LUONG NHAN VIEN", styles["Title"]),
                Paragraph("Ky luong: " + escape(ascii_pdf(ky_luong)), styles["Normal"]), Spacer(1, 8*mm)]
    pairs = [("Ho va ten", nv["Họ và tên"]), ("Email ca nhan", nv["Email cá nhân"]),
             ("Ngay sinh", nv["Ngày sinh"]), ("Chuc vu", nv["Chức vụ"])]
    pairs += [(ascii_pdf(k), f"{float(nv[k]):,.3f} trieu dong") for k in COT_TIEN]
    pairs.append(("So nguoi phu thuoc", str(int(nv["Số người phụ thuộc"]))))
    data = [[Paragraph(escape(ascii_pdf(k)), styles["Normal"]),
             Paragraph(escape(ascii_pdf(v)), styles["Normal"])] for k, v in pairs]
    table = Table(data, colWidths=[70*mm, 100*mm], hAlign="LEFT", repeatRows=0)
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("BACKGROUND", (0, -2), (-1, -2), colors.whitesmoke),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)
    elements.extend([Spacer(1, 8*mm), Paragraph(
        "Luu y: PDF su dung tieng Viet khong dau do khong cai font Unicode.", styles["Italic"])])
    doc.build(elements)
    return buf.getvalue()


def ten_file(s):
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", ascii_pdf(s)).strip("_")[:60] or "nhan_vien"


def tao_zip(df, ky):
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, (_, nv) in enumerate(df.iterrows(), start=1):
            zf.writestr(f"{i:03d}_{ten_file(nv['Họ và tên'])}.pdf", tao_pdf(nv, ky))
    return buf.getvalue()


# ------------------- GIAO DIEN -------------------
logo = Path(__file__).resolve().parent / "LOGO.jpg"
if logo.exists():
    st.image(str(logo), width=170)
st.title("💰 PHẦN MỀM TÍNH TIỀN LƯƠNG")
st.subheader("Song Hào")
st.info("Đơn vị các khoản tiền: TRIỆU ĐỒNG/THÁNG. Thuế và giảm trừ theo công thức mẫu đã cung cấp; cần kiểm tra quy định hiện hành trước khi trả lương thực tế.")
st.caption("PDF không dùng font bổ sung nên sẽ hiển thị tiếng Việt KHÔNG DẤU; Excel và giao diện vẫn đầy đủ dấu.")

if "nhan_vien" not in st.session_state:
    st.session_state.nhan_vien = du_lieu_mac_dinh()
if "so_nv" not in st.session_state:
    st.session_state.so_nv = 10

st.subheader("1. Chọn số lượng người lao động")
so_nv = st.number_input("Số nhân viên", min_value=1, max_value=MAX_NV, step=1, key="so_nv")
ky_luong = st.text_input("Kỳ lương (ví dụ: 10/2026)", value=date.today().strftime("%m/%Y"))

st.subheader("2. Nhập thông tin nhân viên và tiền lương")
st.caption("Có thể dán nhiều ô từ Excel. Ngày sinh nhập dạng DD/MM/YYYY. Email và ngày sinh có thể để trống.")

bang = st.session_state.nhan_vien.iloc[:so_nv].copy()
cfg = {
    "Họ và tên": st.column_config.TextColumn(
        "Họ và tên", required=True, width="medium"
    ),
    "Số điện thoại Zalo": st.column_config.TextColumn(
        "Số điện thoại Zalo",
        width="medium",
        help="Nhập số điện thoại Zalo, ví dụ: 0912345678"
    ),
    "Email cá nhân": st.column_config.TextColumn(
        "Email cá nhân", width="medium"
    ),
    "Ngày sinh": st.column_config.TextColumn("Ngày sinh (DD/MM/YYYY)", width="medium"),
    "Chức vụ": st.column_config.TextColumn("Chức vụ", width="medium"),
    "Số người phụ thuộc": st.column_config.NumberColumn("Số người phụ thuộc", min_value=0, max_value=100, step=1, format="%d"),
}
for c in [
    "Lương cơ bản",
    "Tiền thưởng",
    "Phụ cấp tính BHXH",
    "Phụ cấp không tính BHXH"
]:
    cfg[c] = st.column_config.NumberColumn(c, min_value=0.0, step=0.1, format="%.3f")

bang_nhap = st.data_editor(bang, hide_index=True, use_container_width=True,
                          num_rows="fixed", column_config=cfg, key=f"bang_{so_nv}")
st.session_state.nhan_vien.iloc[:so_nv] = bang_nhap.to_numpy()

if st.button("🧮 TÍNH LƯƠNG TOÀN BỘ NHÂN VIÊN", type="primary", use_container_width=True):
    loi = kiem_tra(bang_nhap)
    if loi:
        st.error("Dữ liệu chưa hợp lệ. Vui lòng sửa trước khi tính:")
        for item in loi[:15]:
            st.write("• " + item)
        if len(loi) > 15:
            st.warning(f"Còn {len(loi)-15} lỗi khác cần sửa.")
    else:
        rows = []
        for _, nv in bang_nhap.iterrows():
            r = nv.to_dict()
            r.update(tinh_luong(nv))
            rows.append(r)
        st.session_state.ket_qua = pd.DataFrame(rows)
        st.session_state.ky_da_tinh = ky_luong
        st.success("Đã tính lương thành công!")

if "ket_qua" in st.session_state:
    df = st.session_state.ket_qua
    ky = st.session_state.get("ky_da_tinh", ky_luong)
    st.subheader("3. Bảng kết quả tính lương")
    a, b, c, d = st.columns(4)
    a.metric("Số nhân viên", len(df))
    b.metric("Tổng thu nhập", f"{df['Tổng thu nhập'].sum():,.3f} triệu")
    c.metric("Tổng thuế TNCN", f"{df['Thuế TNCN'].sum():,.3f} triệu")
    d.metric("Tổng thực lãnh", f"{df['Tiền thực lãnh'].sum():,.3f} triệu")
    st.dataframe(df, hide_index=True, use_container_width=True)

    st.subheader("4. Xuất bảng lương Excel / CSV")
    st.download_button("📥 TẢI FILE EXCEL (.xlsx)", data=tao_excel(df),
                       file_name="bang_luong.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       use_container_width=True)
    st.download_button("📥 TẢI FILE CSV", data=df.to_csv(index=False).encode("utf-8-sig"),
                       file_name="bang_luong.csv", mime="text/csv", use_container_width=True)

    st.subheader("5. Xuất PDF phiếu lương từng nhân viên")
    st.warning("PDF không dấu để tránh phụ thuộc vào font Unicode. Excel vẫn giữ tiếng Việt có dấu.")
    lua_chon = st.selectbox("Chọn nhân viên", options=list(range(len(df))),
                           format_func=lambda i: f"{i+1}. {df.iloc[i]['Họ và tên']}")
    nv_chon = df.iloc[lua_chon]
    st.download_button("📄 TẢI PDF NHÂN VIÊN ĐÃ CHỌN", data=tao_pdf(nv_chon, ky),
                       file_name=f"phieu_luong_{lua_chon+1:03d}_{ten_file(nv_chon['Họ và tên'])}.pdf",
                       mime="application/pdf", use_container_width=True)
    st.download_button("📦 TẢI ZIP PDF CỦA TẤT CẢ NHÂN VIÊN", data=tao_zip(df, ky),
                       file_name="phieu_luong_tat_ca.zip", mime="application/zip", use_container_width=True)
    st.caption("Nếu sửa dữ liệu đầu vào, hãy nhấn TÍNH LƯƠNG lại để cập nhật file xuất.")

with st.expander("📘 Công thức tính lương đang áp dụng"):
    st.markdown("""
    - Tổng thu nhập = Lương cơ bản + Thưởng + Phụ cấp tính BHXH + Phụ cấp không tính BHXH.
    - Bảo hiểm người lao động đóng = (Lương cơ bản + Phụ cấp tính BHXH) × 10,5%.
    - Thu nhập tính thuế = max(0, Tổng thu nhập − Bảo hiểm − 15,5 − 6,2 × Số người phụ thuộc).
    - Thuế TNCN tính theo biểu thuế mẫu 5 bậc trong mã nguồn.
    - Tiền thực lãnh = Tổng thu nhập − Bảo hiểm − Thuế TNCN.

    **Lưu ý:** Các mức và giả định chưa được kiểm chứng cho từng tình huống trả lương.
    """)


# ------------------- HANG DOI GUI ZALO -------------------
if "ket_qua" in st.session_state:
    st.divider()
    st.subheader("📤 Gửi phiếu lương PDF qua Zalo")

    df_zalo = st.session_state.ket_qua
    ky_zalo = st.session_state.get("ky_da_tinh", ky_luong)

    st.caption(
        "Tạo PDF và đưa vào Supabase. "
        "Windows Agent sẽ thực hiện gửi qua Zalo PC sau."
    )

    dau_vao_zalo = (
        st.session_state.nhan_vien.iloc[:so_nv][COT_NHAP]
        .reset_index(drop=True)
    )
    da_tinh_zalo = df_zalo[COT_NHAP].reset_index(drop=True)

    du_lieu_cu = (
        len(dau_vao_zalo) != len(da_tinh_zalo)
        or not dau_vao_zalo.equals(da_tinh_zalo)
        or ky_luong != ky_zalo
    )

    if du_lieu_cu:
        st.warning("Dữ liệu đã thay đổi. Hãy TÍNH LƯƠNG lại.")

    xac_nhan_zalo = st.checkbox(
        "Tôi xác nhận danh sách người nhận và đồng ý tạo "
        "công việc gửi phiếu lương qua Zalo.",
        key="xac_nhan_hang_doi_zalo"
    )

    # Chống tạo lại hàng đợi trong cùng phiên Streamlit
    if "dot_zalo_da_tao" not in st.session_state:
        st.session_state.dot_zalo_da_tao = False

    if st.button(
        "📤 ĐƯA PHIẾU LƯƠNG VÀO HÀNG ĐỢI ZALO",
        disabled=(
            du_lieu_cu
            or not xac_nhan_zalo
            or st.session_state.dot_zalo_da_tao
            or df_zalo.empty
        ),
        use_container_width=True
    ):
        # Khóa tạo lại hàng đợi ngay khi bắt đầu xử lý
        st.session_state.dot_zalo_da_tao = True        
        batch_id = str(uuid.uuid4())
        bao_cao_zalo = []
        so_zalo_da_xu_ly = set()
        thanh_cong = 0
        loi = 0

        progress = st.progress(0)

        for i, (_, nv) in enumerate(df_zalo.iterrows(), start=1):
            ten = str(nv["Họ và tên"]).strip()
            so = chuan_hoa_so_zalo(nv["Số điện thoại Zalo"])
            trang_thai = "Pending"
            ghi_chu = ""

            if not so:
                trang_thai = "Error"
                ghi_chu = "Số điện thoại Zalo không hợp lệ"
                loi += 1

            elif so in so_zalo_da_xu_ly:
                trang_thai = "Error"
                ghi_chu = "Số điện thoại Zalo bị trùng trong đợt"
                loi += 1

            else:
                so_zalo_da_xu_ly.add(so)
                try:
                    pdf = tao_pdf(nv, ky_zalo)

                    pdf_path = (
                        f"{batch_id}/"
                        f"{i:03d}_{uuid.uuid4().hex}.pdf"
                    )

                    tai_pdf_len_supabase(pdf, pdf_path)

                    tao_cong_viec_zalo(
                        batch_id=batch_id,
                        ho_ten=ten,
                        so_zalo=so,
                        ky_luong=ky_zalo,
                        pdf_path=pdf_path
                    )

                    thanh_cong += 1

                except (requests.RequestException, ValueError,
                        RuntimeError, OSError) as exc:
                    ghi_chu = f"{type(exc).__name__}: {str(exc)[:200]}"

                    if str(exc).startswith("Unknown:"):
                        trang_thai = "Unknown"
                    else:
                        trang_thai = "Error"

                    loi += 1

            bao_cao_zalo.append({
                "STT": i,
                "Họ và tên": ten,
                "Số điện thoại Zalo": so or "",
                "Trạng thái": trang_thai,
                "Ghi chú": ghi_chu
            })

            progress.progress(i / len(df_zalo))

        st.session_state["bao_cao_hang_doi_zalo"] = pd.DataFrame(
            bao_cao_zalo
        )
        st.session_state["batch_id_zalo"] = batch_id
                
        # Khóa việc tạo lại đợt sau khi đã thực hiện
        # kể cả khi kết quả có lỗi hoặc chưa xác định
        st.session_state.dot_zalo_da_tao = True
        st.success(
            f"Đã tạo {thanh_cong} công việc Pending. "
            f"Có {loi} nhân viên chưa đưa vào hàng đợi."
        )

    if "bao_cao_hang_doi_zalo" in st.session_state:
        st.markdown("**Báo cáo tạo hàng đợi Zalo**")
        bao_cao = st.session_state["bao_cao_hang_doi_zalo"]
        st.dataframe(
            bao_cao,
            hide_index=True,
            use_container_width=True
        )

        st.download_button(
            "📥 TẢI BÁO CÁO CSV",
            data=bao_cao.to_csv(index=False).encode("utf-8-sig"),
            file_name="bao_cao_hang_doi_zalo.csv",
            mime="text/csv"
        )

# ------------------- GUI PHIEU LUONG QUA GMAIL -------------------
# Bao mat: chi doc thong tin dang nhap tu Streamlit Secrets.
import smtplib
import ssl
import hashlib
import json
from email.message import EmailMessage
from email.utils import formataddr, make_msgid


def email_hop_le(email):
    return bool(re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?)+", str(email).strip()))


def lay_cau_hinh_gmail():
    try:
        gmail = st.secrets["gmail"]
        address = str(gmail["address"]).strip()
        password = str(gmail["app_password"]).replace(" ", "").strip()
        if not email_hop_le(address) or not password:
            return None
        return {"address": address, "app_password": password}
    except (KeyError, FileNotFoundError, AttributeError):
        return None


def tao_email(nv, ky, sender, subject, template):
    ten = str(nv["Họ và tên"])
    email = str(nv["Email cá nhân"]).strip()
    msg = EmailMessage()
    msg["From"] = formataddr(("Bo phan nhan su", sender))
    msg["To"] = email
    msg["Subject"] = subject.replace("{ky_luong}", ky)
    msg["Message-ID"] = make_msgid()
    noi_dung = template.replace("{ho_ten}", ten).replace("{ky_luong}", ky).replace("{chuc_vu}", str(nv["Chức vụ"]))
    msg.set_content(noi_dung, charset="utf-8")
    pdf = tao_pdf(nv, ky)
    filename = f"phieu_luong_{ten_file(ten)}_{ten_file(ky)}.pdf"
    msg.add_attachment(pdf, maintype="application", subtype="pdf", filename=filename)
    return msg


def gui_mot_email(smtp, nv, ky, sender, subject, template):
    msg = tao_email(nv, ky, sender, subject, template)
    smtp.send_message(msg, from_addr=sender, to_addrs=[str(nv["Email cá nhân"]).strip()])


def ma_dot_gui(df, ky, subject, template):
    payload = {
        "ky": ky,
        "subject": subject,
        "template": template,
        "rows": df.fillna("").astype(str).to_dict(orient="records"),
    }
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def tao_bao_cao_gui(rows):
    return pd.DataFrame(rows, columns=["STT", "Họ và tên", "Email cá nhân", "Kỳ lương", "Trạng thái", "Thời gian", "Ghi chú"])


if "ket_qua" in st.session_state:
    st.divider()
    st.subheader("6. Gửi phiếu lương PDF qua Gmail cá nhân")
    st.caption("Chỉ gửi khi bạn bấm xác nhận. Mỗi thư chỉ đính kèm phiếu lương của đúng một nhân viên. PDF hiện là tiếng Việt không dấu.")

    df_mail = st.session_state.ket_qua
    ky_mail = st.session_state.get("ky_da_tinh", ky_luong)
    # Chặn gửi kết quả cũ khi người dùng vừa sửa bảng đầu vào.
    dau_vao_hien_tai = st.session_state.nhan_vien.iloc[:so_nv][COT_NHAP].reset_index(drop=True)
    dau_vao_da_tinh = df_mail[COT_NHAP].reset_index(drop=True)
    du_lieu_da_thay_doi = (
        len(dau_vao_hien_tai) != len(dau_vao_da_tinh)
        or not dau_vao_hien_tai.equals(dau_vao_da_tinh)
        or ky_luong != ky_mail
    )
    if du_lieu_da_thay_doi:
        st.error("Bảng nhân viên hoặc kỳ lương đã thay đổi. Hãy nhấn TÍNH LƯƠNG lại trước khi gửi email.")

    cau_hinh = lay_cau_hinh_gmail()
    if cau_hinh is None:
        st.warning("Chưa cấu hình Gmail. Hãy thêm [gmail] vào Streamlit Secrets (xem hướng dẫn bên dưới). Không nhập mật khẩu trong app.py.")
    else:
        st.success(f"Tài khoản gửi đã cấu hình: {cau_hinh['address']}")

    with st.expander("Cấu hình Gmail an toàn bằng Streamlit Secrets"):
        st.code('[gmail]\naddress = "tenban@gmail.com"\napp_password = "MAT_KHAU_UNG_DUNG_16_KY_TU"', language="toml")
        st.write("Dùng mật khẩu ứng dụng của Google (khi tài khoản hỗ trợ và đã bật xác minh 2 bước). Không dùng mật khẩu Gmail thông thường. Không đưa file secrets.toml lên GitHub hoặc chia sẻ công khai.")
        st.write("Nếu nền tảng không hỗ trợ Secrets, tạo file `.streamlit/secrets.toml` trong môi trường riêng tư và không đưa vào kho mã nguồn công khai.")

    kieu_gui = st.radio("Chế độ gửi", ["Gửi một nhân viên", "Gửi tất cả nhân viên có email"], horizontal=True)
    if kieu_gui == "Gửi một nhân viên":
        chi_so = st.selectbox("Người nhận", list(range(len(df_mail))), format_func=lambda i: f"{i+1}. {df_mail.iloc[i]['Họ và tên']} — {df_mail.iloc[i]['Email cá nhân']}", key="chon_gui_email")
        chi_so_gui = [chi_so]
    else:
        chi_so_gui = list(range(len(df_mail)))

    subject = st.text_input("Tiêu đề email", value="Phiếu lương kỳ {ky_luong}", key="tieu_de_email")
    template = st.text_area("Nội dung email (có thể dùng {ho_ten}, {chuc_vu}, {ky_luong})", value=(
        "Kính gửi {ho_ten},\n\n"
        "Bộ phận Nhân sự gửi Anh/Chị phiếu lương kỳ {ky_luong} đính kèm.\n"
        "Vui lòng kiểm tra và liên hệ bộ phận Nhân sự nếu có thắc mắc.\n\n"
        "Trân trọng,\nBộ phận Nhân sự"
    ), height=180, key="noi_dung_email")

    dot_gui = ma_dot_gui(df_mail, ky_mail, subject, template)
    if st.session_state.get("dot_gui_hien_tai") != dot_gui:
        st.session_state.dot_gui_hien_tai = dot_gui
        st.session_state.da_gui = set()
        st.session_state.bao_cao_gui = []
        st.session_state.xac_nhan_gui = False

    bang_xem_truoc = []
    hop_le = []
    for i in chi_so_gui:
        nv = df_mail.iloc[i]
        email = str(nv["Email cá nhân"]).strip()
        valid = email_hop_le(email)
        da_gui = i in st.session_state.da_gui
        bang_xem_truoc.append({"STT": i+1, "Họ và tên": nv["Họ và tên"], "Email": email,
                              "Kỳ lương": ky_mail, "Trạng thái": "Đã gửi trong đợt này" if da_gui else ("Sẵn sàng" if valid else "Thiếu/sai email")})
        if valid and not da_gui:
            hop_le.append(i)

    st.markdown("**Xem trước danh sách người nhận**")
    st.dataframe(pd.DataFrame(bang_xem_truoc), hide_index=True, use_container_width=True)
    st.caption(f"Sẵn sàng gửi: {len(hop_le)} | Tổng số được chọn: {len(chi_so_gui)}. Không gửi đến địa chỉ thiếu hoặc sai định dạng.")
    if kieu_gui == "Gửi một nhân viên":
        nv_preview = df_mail.iloc[chi_so_gui[0]]
        st.markdown("**Xem trước nội dung thư**")
        st.text("Tiêu đề: " + subject.replace("{ky_luong}", ky_mail))
        st.text(template.replace("{ho_ten}", str(nv_preview["Họ và tên"]))
                .replace("{chuc_vu}", str(nv_preview["Chức vụ"]))
                .replace("{ky_luong}", ky_mail))
        st.download_button("📄 Xem phiếu PDF đính kèm", data=tao_pdf(nv_preview, ky_mail),
                           file_name=f"xem_truoc_{ten_file(nv_preview['Họ và tên'])}.pdf", mime="application/pdf")

    st.checkbox("Tôi đã kiểm tra người nhận, nội dung và đồng ý gửi phiếu lương chứa dữ liệu cá nhân.", key="xac_nhan_gui")
    disabled = du_lieu_da_thay_doi or cau_hinh is None or not hop_le or not st.session_state.xac_nhan_gui or not subject.strip() or not template.strip()
    if st.button(f"📨 XÁC NHẬN GỬI {len(hop_le)} EMAIL", type="primary", disabled=disabled, use_container_width=True):
        # Không tự thử lại các thư lỗi để tránh gửi trùng khi máy chủ đã nhận thư.
        # Mỗi SMTP send_message chỉ có đúng một địa chỉ người nhận.
        thanh_cong = 0
        that_bai = 0
        progress = st.progress(0)
        try:
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context(), timeout=25) as smtp:
                smtp.login(cau_hinh["address"], cau_hinh["app_password"])
                for pos, i in enumerate(hop_le, 1):
                    nv = df_mail.iloc[i]
                    email = str(nv["Email cá nhân"]).strip()
                    trang_thai = "Thất bại"
                    ghi_chu = ""
                    try:
                        gui_mot_email(smtp, nv, ky_mail, cau_hinh["address"], subject, template)
                        st.session_state.da_gui.add(i)
                        trang_thai = "SMTP đã chấp nhận"
                        thanh_cong += 1
                    except (smtplib.SMTPException, OSError, ValueError) as exc:
                        # Không đưa mật khẩu hay nội dung SMTP nhạy cảm vào báo cáo.
                        ghi_chu = type(exc).__name__
                        that_bai += 1
                    st.session_state.bao_cao_gui.append({
                        "STT": i+1, "Họ và tên": nv["Họ và tên"], "Email cá nhân": email,
                        "Kỳ lương": ky_mail, "Trạng thái": trang_thai,
                        "Thời gian": datetime.now().strftime("%d/%m/%Y %H:%M:%S"), "Ghi chú": ghi_chu
                    })
                    progress.progress(pos / len(hop_le))
        except (smtplib.SMTPException, OSError, ValueError) as exc:
            st.error(f"Không thể kết nối/đăng nhập Gmail ({type(exc).__name__}). Kiểm tra Gmail App Password và quyền kết nối SMTP của nền tảng.")
            st.session_state["da_gui_email"] = True
        if thanh_cong:
            st.success(f"Máy chủ SMTP đã chấp nhận {thanh_cong} email. Điều này không bảo đảm thư đã vào hộp thư đến.")
        if that_bai:
            st.warning(f"Có {that_bai} email không gửi được; xem báo cáo phía dưới.")

    if st.session_state.get("bao_cao_gui"):
        st.markdown("**Báo cáo kết quả gửi email trong đợt này**")
        bao_cao = tao_bao_cao_gui(st.session_state.bao_cao_gui)
        st.dataframe(bao_cao, hide_index=True, use_container_width=True)
        st.download_button("📥 Tải báo cáo CSV", bao_cao.to_csv(index=False).encode("utf-8-sig"),
                           file_name="bao_cao_gui_phieu_luong.csv", mime="text/csv")
        st.caption("Báo cáo chỉ tồn tại trong phiên làm việc này. Không tự gửi lại email đã được SMTP chấp nhận trong cùng đợt.")


# KIỂM TRA KẾT NỐI SUPABASE
with st.expander("🔗 Kiểm tra kết nối Supabase"):
    try:
        import requests

        sb = st.secrets["supabase"]
        url = str(sb["url"]).strip().rstrip("/")
        key = str(sb["secret_key"]).strip()

        response = requests.get(
            f"{url}/rest/v1/zalo_salary_jobs",
            headers={
                "apikey": key,
                "Authorization": f"Bearer {key}"
            },
            params={"select": "id", "limit": "1"},
            timeout=15
        )

        if response.status_code == 200:
            st.success("OK - Đã kết nối thành công với Supabase")
        else:
            st.error(
                f"Error - Supabase trả về HTTP {response.status_code}"
            )

    except Exception as e:
        st.error(f"Error - {type(e).__name__}")
